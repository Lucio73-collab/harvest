"""Export items to JSON, JSONL, CSV or Excel."""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Any

FORMATS = ("json", "jsonl", "csv", "xlsx")


def _columns(items: list[dict[str, Any]], include_internal: bool) -> list[str]:
    cols: list[str] = []
    for it in items:
        for k in it:
            if k not in cols and (include_internal or not k.startswith("_")):
                cols.append(k)
    return cols


def _cell(v: Any) -> Any:
    if isinstance(v, (dict, list)):
        if isinstance(v, list) and all(isinstance(x, (str, int, float)) for x in v):
            return " | ".join(map(str, v))
        return json.dumps(v, ensure_ascii=False)
    return v


def to_bytes(items: list[dict[str, Any]], fmt: str, include_internal: bool = False) -> bytes:
    if not include_internal:
        items = [{k: v for k, v in i.items() if not k.startswith("_") or k == "_url"} for i in items]
    if fmt == "json":
        return json.dumps(items, ensure_ascii=False, indent=2, default=str).encode()
    if fmt == "jsonl":
        return "\n".join(json.dumps(i, ensure_ascii=False, default=str) for i in items).encode()
    cols = _columns(items, include_internal=True)
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(cols)
        for i in items:
            w.writerow([_cell(i.get(c)) for c in cols])
        return ("﻿" + buf.getvalue()).encode()   # BOM so Excel opens UTF-8 correctly
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        ws = wb.active
        ws.title = "items"
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="111111")
            c.fill = PatternFill("solid", fgColor="A6E3A1")
        for i in items:
            row = []
            for c in cols:
                v = _cell(i.get(c))
                if isinstance(v, str) and len(v) > 32000:   # Excel cell limit
                    v = v[:32000] + "..."
                row.append(v)
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for idx, c in enumerate(cols, 1):
            width = max([len(str(c))] + [len(str(_cell(i.get(c)) or "")) for i in items[:200]])
            ws.column_dimensions[get_column_letter(idx)].width = max(10, min(60, width + 2))
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    raise ValueError(f"unknown format {fmt!r}; use one of {FORMATS}")


def write(items: list[dict[str, Any]], path: str | Path, fmt: str | None = None) -> Path:
    p = Path(path)
    fmt = fmt or p.suffix.lstrip(".").lower() or "json"
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(to_bytes(items, fmt))
    return p
